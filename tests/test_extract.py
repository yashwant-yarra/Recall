import io
import pytest
from docx import Document
from recall.extract import extract_text

def test_txt():
    assert extract_text("héllo".encode(), "text/plain") == ["héllo"]
def test_docx():
    doc = Document()
    doc.add_paragraph("first")
    doc.add_paragraph("second")
    buf = io.BytesIO()
    doc.save(buf)
    assert extract_text(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document") == ["first\nsecond"]

def test_unsupported():
    with pytest.raises(ValueError):
        extract_text(b"x", "image/png")